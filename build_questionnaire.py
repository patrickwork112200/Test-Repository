import math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/home/user/Test-Repository/DASH_Hypercare_Discovery_Questionnaire.xlsx"
VERSION = "v2.0"
DATE = "17 June 2026"

# ---------------------------------------------------------------- content
SECTIONS = [
    ("Engagement Scope & Support Activities", [
        "What services are in scope (e.g., L1 / L2 / L3 support, incident, change, release, monitoring)?",
        "What is explicitly out of scope?",
        "What is the ownership split across application, infrastructure, and integrations?",
        "Is support reactive only, or does it also include proactive services (e.g., monitoring, problem management, continuous improvement)?",
        "What activities are expected under support (e.g., incident management, bug fixing, monitoring, deployment support, performance optimization, user support, minor enhancements)?",
        "Are there any special support requirements (e.g., data-loading, month-end or quarter-end support)?",
    ]),
    ("Application & Technical Landscape", [
        "What applications and interfaces are in scope, including their environments?",
        "For each in-scope application, please provide the technology stack and a business-criticality rating (1 = most critical … 5 = least critical).",
        "How are the applications classified by tier (Tier 1 / 2 / 3)?",
        "Which business functions and processes does each application support?",
        "What is the current OutSystems version / platform in use?",
        "Are the applications hosted on cloud or on-premises?",
        "Are there any third-party integrations? If so, please describe them.",
        "Are any components or customizations classified as proprietary?",
        "How many active users does the application support?",
        "What is the approximate transaction / load volume?",
        "Are there any known unstable systems or high-defect modules?",
        "What is the current level of technical debt?",
    ]),
    ("Operational Volumes & Demand Profile", [
        "What is the monthly ticket volume by type and severity (incidents, service requests, changes)?",
        "What is the hourly / daily demand distribution (peak vs off-peak)?",
        "What percentage of tickets are repetitive vs complex?",
        "What percentage of tickets require escalation (L1 → L2 → L3)?",
        "What are the top recurring issues and root causes?",
        "Are there seasonal spikes or predictable peaks affecting demand?",
        "What is the current automation coverage (% of tickets auto-resolved)?",
    ]),
    ("Performance Baselines & Routine Operations", [
        "What is the average handling time (AHT) per ticket type?",
        "What are the current performance baselines (MTTR, FCR, backlog, SLA compliance)?",
        "What daily, weekly, and monthly activities are required to keep the application operational (e.g., housekeeping, monitoring)?",
        "What is the business impact of downtime per application tier?",
    ]),
    ("Current Operating Model, Tooling & Documentation", [
        "What is the current team model (roles, FTEs, location, vendor mix, skill coverage)?",
        "What tools are used (e.g., ITSM, monitoring, CI/CD, reporting)?",
        "What technical and functional documentation exists (e.g., runbooks, SOPs, knowledge-base articles, architecture diagrams)? Please share copies or indicate where they can be accessed.",
    ]),
    ("Change & Release Management", [
        "How many releases occur per month (planned vs emergency)?",
        "What is the average effort per change?",
        "What are the deployment windows (e.g., weekends, after-hours)?",
        "What percentage of effort is spent on support vs enhancements?",
    ]),
    ("Service Levels & KPIs", [
        "What are the required response and resolution SLAs by severity (P1–P4)?",
        "What are the availability targets per application tier?",
        "Which KPIs matter most (e.g., stability, speed, CSAT, cost)?",
        "What are the baseline vs target improvements at 3, 6, and 12 months?",
        "Is there an SLA ramp-up period before full SLAs take effect?",
    ]),
    ("Commercial Model", [
        "What pricing model is preferred (e.g., fixed, per-ticket, outcome-based, hybrid)?",
        "What is the expected budget baseline or cost-reduction target (%)?",
        "How should demand variability be handled (e.g., volume bands, flex capacity)?",
        "What service-credit or penalty expectations exist?",
        "Are there incentives for performance or automation (e.g., gain-share)?",
    ]),
    ("Governance & Reporting", [
        "What governance cadence is expected (e.g., daily operations, MBR, QBR)?",
        "Who are the key stakeholders, and what are the escalation paths?",
        "Is this a multi-vendor environment (SIAM)?",
        "What reporting and transparency are required?",
    ]),
    ("Risk, Security & Compliance", [
        "What regulatory or compliance requirements apply?",
        "What security controls are required?",
        "What are the BCP / DR expectations?",
        "Are there any data-residency restrictions?",
    ]),
]

# ---------------------------------------------------------------- palette
HEADER_FILL = "14365C"
SECTION_COLORS = [
    ("14365C", "FFFFFF"), ("1F4E78", "FFFFFF"), ("2A5E8C", "FFFFFF"),
    ("356E9E", "FFFFFF"), ("4A7FAE", "FFFFFF"), ("5E90BE", "FFFFFF"),
    ("7AA7CE", "1F2A37"), ("9BBFDD", "1F2A37"), ("BBD4E9", "1F2A37"),
    ("D6E5F2", "1F2A37"),
]
TITLE_FILL = "14365C"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WHITE = "FFFFFF"

wb = Workbook()

# ================================================================ INSTRUCTIONS
ins = wb.active
ins.title = "Instructions"
ins.sheet_view.showGridLines = False
ins.column_dimensions["A"].width = 3
ins.column_dimensions["B"].width = 30
ins.column_dimensions["C"].width = 95

def ins_title(r, text, size, fill=None, color="FFFFFF", bold=True):
    ins.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    c = ins.cell(row=r, column=2, value=text)
    c.font = Font(name="Calibri", size=size, bold=bold, color=color)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    if fill:
        for col in (2, 3):
            ins.cell(row=r, column=col).fill = PatternFill("solid", fgColor=fill)

def ins_kv(r, k, v):
    a = ins.cell(row=r, column=2, value=k)
    a.font = Font(bold=True, size=11, color="14365C")
    a.alignment = Alignment(vertical="top", wrap_text=True)
    b = ins.cell(row=r, column=3, value=v)
    b.font = Font(size=11)
    b.alignment = Alignment(vertical="top", wrap_text=True)

r = 1
ins.row_dimensions[r].height = 36
ins_title(r, "DASH – Managed Services & Hypercare", 18, fill=TITLE_FILL)
r += 1
ins.row_dimensions[r].height = 24
ins_title(r, "Discovery Questionnaire", 13, fill="1F4E78")
r += 2

ins_kv(r, "Prepared for", "AXA"); r += 1
ins_kv(r, "Version", f"{VERSION}  ·  {DATE}"); r += 1
ins_kv(r, "Return to", "[Service delivery lead – name / email]"); r += 1
ins_kv(r, "Response due", "[Response due date]"); r += 2

ins_title(r, "Purpose", 12, color="14365C"); r += 1
ins.cell(row=r, column=3,
    value=("This questionnaire gathers the information needed to design and "
           "transition the managed-services and hypercare model for the DASH "
           "application estate. Your responses establish the operational "
           "baseline, scope, service levels, and commercial expectations that "
           "shape the proposal and transition plan.")
    ).alignment = Alignment(wrap_text=True, vertical="top")
ins.cell(row=r, column=3).font = Font(size=11)
ins.row_dimensions[r].height = 60
r += 2

ins_title(r, "How to complete", 12, color="14365C"); r += 1
steps = [
    "Work through the “Questionnaire” tab. Questions are grouped into 10 sections and numbered for easy reference (e.g., 2.3).",
    "Enter your answer in the “AXA Response” column. Use the “Comments / Supporting Documents” column to add context or name any attachments.",
    "Where a question does not apply, enter “N/A” and a brief reason rather than leaving it blank.",
    "If a precise figure is unavailable, provide your best estimate and flag it as an estimate.",
    "Attach supporting documents (runbooks, architecture diagrams, reports) separately and reference them by file name in the Comments column.",
    "Use the column filters (drop-down arrows on the header row) to focus on a single section, or assign sections to different teams.",
]
for s in steps:
    c = ins.cell(row=r, column=3, value="•  " + s)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.font = Font(size=11)
    ins.row_dimensions[r].height = 30
    r += 1
r += 1

ins_title(r, "Column legend", 12, color="14365C"); r += 1
legend = [
    ("Ref", "Unique question reference (section.question)."),
    ("Section", "Topic grouping for the question."),
    ("Question", "The information being requested."),
    ("AXA Response", "Your answer – please complete."),
    ("Comments / Supporting Documents", "Context, assumptions, or referenced attachments."),
]
for k, v in legend:
    ins_kv(r, k, v); ins.row_dimensions[r].height = 22; r += 1
r += 1

ins_title(r, "Section index", 12, color="14365C"); r += 1
total_q = 0
for i, (name, qs) in enumerate(SECTIONS, start=1):
    total_q += len(qs)
    a = ins.cell(row=r, column=2, value=f"Section {i}")
    a.font = Font(bold=True, size=11, color="14365C")
    a.alignment = Alignment(vertical="center")
    b = ins.cell(row=r, column=3, value=f"{name}  —  {len(qs)} questions")
    b.font = Font(size=11)
    b.alignment = Alignment(vertical="center")
    r += 1
a = ins.cell(row=r, column=2, value="Total")
a.font = Font(bold=True, size=11, color="14365C")
ins.cell(row=r, column=3, value=f"{total_q} questions across {len(SECTIONS)} sections").font = Font(bold=True, size=11)
r += 2

ins_title(r, "Confidentiality", 12, color="14365C"); r += 1
ins.cell(row=r, column=3,
    value=("The information shared in this questionnaire is treated as "
           "confidential and used solely for the purpose of scoping and "
           "delivering the DASH managed-services engagement.")
    ).alignment = Alignment(wrap_text=True, vertical="top")
ins.cell(row=r, column=3).font = Font(size=11, italic=True)
ins.row_dimensions[r].height = 45

# ================================================================ QUESTIONNAIRE
qs = wb.create_sheet("Questionnaire")
qs.sheet_view.showGridLines = False
widths = {"A": 8, "B": 32, "C": 82, "D": 50, "E": 40}
for col, w in widths.items():
    qs.column_dimensions[col].width = w

NCOLS = 5
# Title banner
qs.merge_cells("A1:E1")
t = qs.cell(row=1, column=1, value="DASH – Managed Services & Hypercare  |  Discovery Questionnaire")
t.font = Font(size=15, bold=True, color="FFFFFF")
t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
qs.row_dimensions[1].height = 30
qs.merge_cells("A2:E2")
sub = qs.cell(row=2, column=1,
    value="Prepared for AXA  ·  Please complete the “AXA Response” and “Comments / Supporting Documents” columns.")
sub.font = Font(size=10, italic=True, color="FFFFFF")
sub.alignment = Alignment(horizontal="left", vertical="center", indent=1)
qs.row_dimensions[2].height = 20
for r_ in (1, 2):
    for c_ in range(1, NCOLS + 1):
        qs.cell(row=r_, column=c_).fill = PatternFill("solid", fgColor=TITLE_FILL if r_ == 1 else "1F4E78")

HEADER_ROW = 4
headers = ["Ref", "Section", "Question", "AXA Response", "Comments / Supporting Documents"]
for j, h in enumerate(headers, start=1):
    c = qs.cell(row=HEADER_ROW, column=j, value=h)
    c.font = Font(bold=True, size=11, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=HEADER_FILL)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
qs.row_dimensions[HEADER_ROW].height = 22

def est_height(text, width_chars):
    lines = 0
    for part in str(text).split("\n"):
        lines += max(1, math.ceil(len(part) / width_chars))
    return max(28, lines * 15 + 8)

row = HEADER_ROW + 1
for i, (name, questions) in enumerate(SECTIONS, start=1):
    fill_hex, font_hex = SECTION_COLORS[i - 1]
    for k, q in enumerate(questions, start=1):
        ref = f"{i}.{k}"
        # Ref
        cA = qs.cell(row=row, column=1, value=ref)
        cA.font = Font(bold=True, size=11, color=font_hex)
        cA.fill = PatternFill("solid", fgColor=fill_hex)
        cA.alignment = Alignment(horizontal="center", vertical="top")
        # Section
        cB = qs.cell(row=row, column=2, value=name)
        cB.font = Font(bold=(k == 1), size=11, color=font_hex)
        cB.fill = PatternFill("solid", fgColor=fill_hex)
        cB.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
        # Question
        cC = qs.cell(row=row, column=3, value=q)
        cC.font = Font(size=11, color="1F2A37")
        cC.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
        # Response + Comments (input cells)
        cD = qs.cell(row=row, column=4)
        cE = qs.cell(row=row, column=5)
        for cc in (cD, cE):
            cc.fill = PatternFill("solid", fgColor="FFFFFF")
            cc.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
            cc.font = Font(size=11)
        for cc in (cA, cB, cC, cD, cE):
            cc.border = BORDER
        qs.row_dimensions[row].height = est_height(q, 80)
        row += 1

last = row - 1
qs.auto_filter.ref = f"A{HEADER_ROW}:E{last}"
qs.freeze_panes = "A5"
qs.sheet_view.zoomScale = 100

# Page setup for clean printing
qs.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
qs.page_setup.orientation = "landscape"
qs.page_setup.fitToWidth = 1
qs.page_setup.fitToHeight = 0
qs.sheet_properties.pageSetUpPr.fitToPage = True

wb.save(OUT)
print("Saved:", OUT)
print("Total questions:", total_q, "Sections:", len(SECTIONS))
